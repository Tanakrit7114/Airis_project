from __future__ import annotations

import hashlib
import json
import mimetypes
import re
import sqlite3
import uuid
from pathlib import Path
from typing import Any

from app.config import MEMORY_DB
from app.documents.ocr import MAX_FILE_BYTES, process_document

CHUNK_SIZE = 2400
CHUNK_OVERLAP = 250
MAX_SEARCH_RESULTS = 8
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".xml", ".html", ".htm",
    ".css", ".js", ".jsx", ".ts", ".tsx", ".py", ".java", ".c", ".cpp", ".h",
    ".hpp", ".cs", ".go", ".rs", ".swift", ".kt", ".kts", ".rb", ".php", ".sql",
    ".yml", ".yaml", ".toml", ".ini", ".cfg", ".conf", ".log", ".tex",
}


def _safe_name(filename: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", Path(filename).name)[:180] or "file"


def _chunks(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_SIZE)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(0, end - CHUNK_OVERLAP)
    return chunks


def _extract_text(filename: str, content: bytes) -> tuple[str, str, int | None, int]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf" or suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        result = process_document(filename, content, max_text_chars=None)
        return str(result.get("text", "")), str(result.get("engine", "ocr")), result.get("pages"), int(result.get("ocr_pages", 0))

    if suffix in TEXT_EXTENSIONS:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("utf-8", errors="replace")
        return text, "text-decoder", 1, 0

    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to index DOCX files") from exc
        from io import BytesIO
        doc = Document(BytesIO(content))
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            value = paragraph.text.strip()
            if value:
                parts.append(value)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
        return "\n".join(parts), "python-docx", 1, 0

    return "", "archive-only", None, 0


class DocumentStore:
    """Persistent document archive + lightweight searchable text index.

    Originals live under data/documents/ while metadata/chunks live in SQLite.
    Files are deduplicated by SHA-256, so the same upload is not stored twice.
    """

    def __init__(self, db_path: str = MEMORY_DB, storage_dir: str = "data/documents"):
        self.db_path = db_path
        self.storage_dir = Path(storage_dir)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        if db_path != ":memory:":
            Path(db_path).parent.mkdir(parents=True, exist_ok=True)
        self._init()

    def connect(self):
        con = sqlite3.connect(self.db_path, timeout=30)
        con.row_factory = sqlite3.Row
        return con

    def _init(self) -> None:
        with self.connect() as con:
            con.execute("PRAGMA foreign_keys=ON")
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY,
                    filename TEXT NOT NULL,
                    sha256 TEXT NOT NULL UNIQUE,
                    mime_type TEXT NOT NULL DEFAULT 'application/octet-stream',
                    size_bytes INTEGER NOT NULL,
                    storage_path TEXT NOT NULL,
                    extracted_text TEXT NOT NULL DEFAULT '',
                    extraction_engine TEXT NOT NULL DEFAULT 'archive-only',
                    pages INTEGER,
                    ocr_pages INTEGER NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    last_accessed_at TEXT
                )
                """
            )
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS document_chunks (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL,
                    content TEXT NOT NULL,
                    FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE,
                    UNIQUE(document_id, chunk_index)
                )
                """
            )
            con.execute("CREATE INDEX IF NOT EXISTS idx_documents_created ON documents(created_at DESC)")
            con.execute("CREATE INDEX IF NOT EXISTS idx_chunks_document ON document_chunks(document_id, chunk_index)")
            # FTS5 is optional; SQLite builds used on macOS normally include it.
            try:
                con.execute(
                    """
                    CREATE VIRTUAL TABLE IF NOT EXISTS document_chunks_fts
                    USING fts5(content, document_id UNINDEXED, chunk_id UNINDEXED)
                    """
                )
            except sqlite3.OperationalError:
                pass
            con.commit()

    def save(self, filename: str, content: bytes) -> dict[str, Any]:
        if len(content) > MAX_FILE_BYTES:
            raise ValueError("ไฟล์ใหญ่เกิน 25 MB")
        sha256 = hashlib.sha256(content).hexdigest()
        with self.connect() as con:
            existing = con.execute("SELECT * FROM documents WHERE sha256 = ?", (sha256,)).fetchone()
            if existing:
                con.execute("UPDATE documents SET last_accessed_at = CURRENT_TIMESTAMP WHERE id = ?", (existing["id"],))
                con.commit()
                return self._row_dict(existing, reused=True)

        extracted, engine, pages, ocr_pages = _extract_text(filename, content)
        document_id = str(uuid.uuid4())
        suffix = Path(filename).suffix.lower()
        stored_name = f"{sha256}{suffix or '.bin'}"
        storage_path = self.storage_dir / stored_name
        storage_path.write_bytes(content)
        mime_type = mimetypes.guess_type(filename)[0] or "application/octet-stream"

        with self.connect() as con:
            con.execute(
                """
                INSERT INTO documents
                (id, filename, sha256, mime_type, size_bytes, storage_path, extracted_text,
                 extraction_engine, pages, ocr_pages)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (document_id, Path(filename).name, sha256, mime_type, len(content), str(storage_path),
                 extracted, engine, pages, ocr_pages),
            )
            for idx, chunk in enumerate(_chunks(extracted)):
                cur = con.execute(
                    "INSERT INTO document_chunks(document_id, chunk_index, content) VALUES (?, ?, ?)",
                    (document_id, idx, chunk),
                )
                chunk_id = int(cur.lastrowid)
                try:
                    con.execute(
                        "INSERT INTO document_chunks_fts(rowid, content, document_id, chunk_id) VALUES (?, ?, ?, ?)",
                        (chunk_id, chunk, document_id, chunk_id),
                    )
                except sqlite3.OperationalError:
                    pass
            con.commit()

        return {
            "id": document_id,
            "filename": Path(filename).name,
            "download_url": f"/api/documents/{document_id}/download",
            "sha256": sha256,
            "mime_type": mime_type,
            "size_bytes": len(content),
            "pages": pages,
            "ocr_pages": ocr_pages,
            "engine": engine,
            "characters": len(extracted),
            "searchable": bool(extracted.strip()),
            "reused": False,
        }

    @staticmethod
    def _row_dict(row: sqlite3.Row, reused: bool = False) -> dict[str, Any]:
        document_id = row["id"]
        return {
            "id": document_id,
            "filename": row["filename"],
            # Keep download links in every document response so clients do
            # not need to reconstruct API paths (and so list views can offer
            # a download action without fetching the full document first).
            "download_url": f"/api/documents/{document_id}/download",
            "sha256": row["sha256"],
            "mime_type": row["mime_type"],
            "size_bytes": row["size_bytes"],
            "pages": row["pages"],
            "ocr_pages": row["ocr_pages"],
            "engine": row["extraction_engine"],
            "characters": len(row["extracted_text"] or ""),
            "searchable": bool((row["extracted_text"] or "").strip()),
            "created_at": row["created_at"],
            "reused": reused,
        }

    def get(self, document_id: str) -> dict[str, Any] | None:
        with self.connect() as con:
            row = con.execute("SELECT * FROM documents WHERE id = ?", (document_id,)).fetchone()
            if not row:
                return None
            con.execute("UPDATE documents SET last_accessed_at = CURRENT_TIMESTAMP WHERE id = ?", (document_id,))
            con.commit()
            result = self._row_dict(row)
            result["text"] = row["extracted_text"] or ""
            result["storage_path"] = row["storage_path"]
            return result

    def list(self, limit: int = 100) -> list[dict[str, Any]]:
        with self.connect() as con:
            rows = con.execute("SELECT * FROM documents ORDER BY created_at DESC LIMIT ?", (max(1, min(limit, 500)),)).fetchall()
        return [self._row_dict(row) for row in rows]

    def search(self, query: str, limit: int = MAX_SEARCH_RESULTS) -> list[dict[str, Any]]:
        query = (query or "").strip()
        if not query:
            return []
        limit = max(1, min(limit, MAX_SEARCH_RESULTS))
        with self.connect() as con:
            rows: list[sqlite3.Row] = []
            # FTS first; useful for English/space-delimited queries.
            try:
                terms = re.sub(r"[^\w\u0E00-\u0E7F\- ]+", " ", query).strip()
                if terms:
                    fts_query = " OR ".join(part for part in terms.split() if part)
                    rows = con.execute(
                        """
                        SELECT d.*, c.id AS chunk_id, c.chunk_index, c.content,
                               bm25(document_chunks_fts) AS rank
                        FROM document_chunks_fts
                        JOIN document_chunks c ON c.id = document_chunks_fts.chunk_id
                        JOIN documents d ON d.id = c.document_id
                        WHERE document_chunks_fts MATCH ?
                        ORDER BY rank
                        LIMIT ?
                        """,
                        (fts_query, limit),
                    ).fetchall()
            except sqlite3.OperationalError:
                rows = []

            if not rows:
                # LIKE fallback works with Thai and filenames and needs no tokenizer.
                needle = f"%{query}%"
                rows = con.execute(
                    """
                    SELECT d.*, c.id AS chunk_id, c.chunk_index, c.content, 1.0 AS rank
                    FROM document_chunks c
                    JOIN documents d ON d.id = c.document_id
                    WHERE c.content LIKE ? OR d.filename LIKE ?
                    ORDER BY d.created_at DESC, c.chunk_index ASC
                    LIMIT ?
                    """,
                    (needle, needle, limit),
                ).fetchall()

            if not rows:
                # Token fallback: score by how many query terms occur in each chunk.
                tokens = [t for t in re.split(r"\s+", query.lower()) if len(t) >= 2]
                chunks = con.execute(
                    """
                    SELECT d.*, c.id AS chunk_id, c.chunk_index, c.content
                    FROM document_chunks c
                    JOIN documents d ON d.id = c.document_id
                    ORDER BY d.created_at DESC, c.chunk_index ASC
                    LIMIT 1000
                    """
                ).fetchall()
                scored = []
                for row in chunks:
                    hay = row["content"].lower()
                    score = sum(hay.count(token) for token in tokens)
                    if score:
                        scored.append((score, row))
                scored.sort(key=lambda x: (-x[0], x[1]["created_at"], x[1]["chunk_index"]))
                rows = [row for _, row in scored[:limit]]

            result: list[dict[str, Any]] = []
            for row in rows:
                result.append({
                    "document_id": row["id"],
                    "filename": row["filename"],
                    "chunk_index": row["chunk_index"],
                    "content": row["content"],
                    "created_at": row["created_at"],
                })
            return result

    def context(self, query: str, limit: int = 5) -> tuple[str, list[dict[str, Any]]]:
        hits = self.search(query, limit=limit)
        if not hits:
            return "", []
        sections = [
            "Persistent file memory (documents uploaded previously):"
        ]
        for hit in hits:
            sections.append(
                f"\n[File: {hit['filename']} | chunk {hit['chunk_index']}]\n{hit['content']}"
            )
        return "\n".join(sections), hits
